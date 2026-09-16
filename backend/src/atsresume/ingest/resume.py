"""Resume ingest: bytes in, reading-order text plus how it was read out."""

from __future__ import annotations

import io
import logging

from ..config import get_settings
from ..models import SourceDocument, SourceKind
from .pdf_layout import PdfReadError, read_pdf

logger = logging.getLogger(__name__)

SUPPORTED = {"pdf", "docx", "txt", "md", "markdown"}


class ResumeIngestError(RuntimeError):
    """Shown to the candidate, so the message has to say what to do next."""


def _extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _read_docx(data: bytes) -> str:
    """Text from a .docx, including table cells.

    Table cells matter: a large share of resumes lay out dates and employers in
    an invisible table, and a reader that walks only ``document.paragraphs``
    silently drops every one of them.
    """
    try:
        import docx  # imported lazily; python-docx is slow to import
    except ImportError as exc:  # pragma: no cover
        raise ResumeIngestError("python-docx is not installed.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeIngestError(
            f"That .docx could not be opened. It may be corrupt, or a .doc renamed to .docx ({exc})."
        ) from exc

    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # A merged cell repeats its text across the row; collapse duplicates.
            deduped: list[str] = []
            for cell in cells:
                if cell and (not deduped or deduped[-1] != cell):
                    deduped.append(cell)
            if deduped:
                parts.append("  ".join(deduped))

    # Headers and footers carry contact details often enough to be worth reading.
    for section in document.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                if p.text.strip():
                    parts.append(p.text)

    return "\n".join(parts)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ").replace("﻿", "")
    lines = [line.rstrip() for line in text.split("\n")]
    out: list[str] = []
    blank = 0
    for line in lines:
        if line.strip():
            blank = 0
            out.append(line)
        else:
            blank += 1
            if blank <= 1:
                out.append("")
    return "\n".join(out).strip()


def ingest_resume(filename: str, data: bytes) -> SourceDocument:
    settings = get_settings()

    if not data:
        raise ResumeIngestError("That file is empty.")
    if len(data) > settings.max_upload_bytes:
        mb = settings.max_upload_bytes // (1024 * 1024)
        raise ResumeIngestError(f"That file is larger than {mb} MB. Upload a resume, not a portfolio.")

    ext = _extension(filename)
    if ext not in SUPPORTED:
        raise ResumeIngestError(
            f"Unsupported format '.{ext or '?'}'. Upload a PDF, DOCX, TXT or Markdown file."
        )

    notes: list[str] = []

    if ext == "pdf":
        try:
            layout = read_pdf(data)
        except PdfReadError as exc:
            raise ResumeIngestError(str(exc)) from exc

        if layout.style.column_count > 1:
            notes.append(
                "This resume is laid out in two columns. Each column was read separately — "
                "without that the sidebar interleaves with the body and every extracted fact "
                "is scrambled. Many ATS parsers do not do this, which is a good reason to "
                "send the single-column version."
            )
        if layout.page_count > 2:
            notes.append(
                f"{layout.page_count} pages. Most recruiters read the first page; the tailored "
                "version targets one to two."
            )
        doc = SourceDocument(
            kind=SourceKind.PDF,
            raw_text=layout.text,
            page_count=layout.page_count,
            style=layout.style,
            notes=notes,
        )

    elif ext == "docx":
        doc = SourceDocument(kind=SourceKind.DOCX, raw_text=_read_docx(data), notes=notes)

    else:
        try:
            decoded = data.decode("utf-8")
        except UnicodeDecodeError:
            decoded = data.decode("latin-1", errors="replace")
            notes.append("That file was not valid UTF-8; it was decoded as Latin-1.")
        doc = SourceDocument(kind=SourceKind.TEXT, raw_text=decoded, notes=notes)

    doc.raw_text = _normalise(doc.raw_text)

    if len(doc.raw_text) < settings.min_resume_chars:
        raise ResumeIngestError(
            "Almost no text came out of that file. If it is a scanned or image-based PDF, "
            "export a text PDF or paste the resume as text instead."
        )
    if len(doc.raw_text) > settings.max_resume_chars:
        doc.raw_text = doc.raw_text[: settings.max_resume_chars]
        doc.notes.append(
            f"The resume was truncated to {settings.max_resume_chars} characters. "
            "Only the beginning was analysed."
        )

    return doc


def ingest_resume_text(text: str) -> SourceDocument:
    """For pasted text, which has no format to read."""
    settings = get_settings()
    cleaned = _normalise(text)
    if len(cleaned) < settings.min_resume_chars:
        raise ResumeIngestError(
            "Paste at least a few hundred characters of your resume, or upload the file."
        )
    return SourceDocument(
        kind=SourceKind.TEXT,
        raw_text=cleaned[: settings.max_resume_chars],
        notes=["Pasted text carries no formatting, so only the ATS layout is produced."],
    )
